"""Build the 3-slide Nature Opportunity deck on the EY template.

Usage:
    python build_deck.py --template EY_Template.pptx --workbook <Client>_<Country>_Nature_Mapping.xlsx --out <Client>_Nature_Priorities.pptx

Inputs:
    --template  The EY .pptx with a cover slide (title placeholder) and a content slide whose
                title text box reads "Insert Slide Title" and subtitle reads "Client Name".
    --workbook  A completed Nature_Opportunity_Master_Template.xlsx, saved in Excel so formula
                results are stored (Scoring, Taxonomy and Slide_Content tabs are read).

Requires: python-pptx, openpyxl.
"""
import argparse
import copy
import sys

import openpyxl
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

LIGHT, BOLD_FONT = "EYInterstate Light", "EYInterstate"
YELLOW, WHITE, GREY, CARD, LINE = "FFE600", "F2F2F2", "BFBFC7", "2E2E38", "3A3A46"
ROW1, ROW2, DARK, MEDIUM = "16161D", "22222B", "1A1A24", "5F5F6B"
RULE = ["FFE600", "F4B183", "B565D9", "4E95D9"]
LEFT, RIGHT = 0.58, 12.04  # content edges, aligned to the template's coloured rule
WIDTH = RIGHT - LEFT


# ---------------------------------------------------------------- workbook
def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    content = {}
    for key, value, *_ in wb["Slide_Content"].iter_rows(min_row=5, values_only=True):
        if key:
            content[str(key).strip()] = "" if value is None else str(value).strip()
    tax = {r[2]: r for r in wb["Taxonomy"].iter_rows(min_row=5, values_only=True) if r[2]}
    rows = []
    for r in wb["Scoring"].iter_rows(min_row=5, values_only=True):
        if not r[1]:
            continue
        if r[9] in (None, ""):
            sys.exit("Scoring has no calculated scores. Enter C1, C2, C4 for all 15 sub-themes, "
                     "open the workbook in Excel and save it, then run again.")
        t = tax.get(r[1])
        nbsap = "– | 0" if not t or not t[6] else (
            f"${t[7]:.1f}bn | {int(t[6])}" if t[7] >= 0.1 else f"<$0.1bn | {int(t[6])}")
        rows.append(dict(rank=r[0], id=r[1], name=r[3], score=r[9], tier=r[10],
                         evidence=r[19] or "", exposure=r[20] or "", nbsap=nbsap))
    rows.sort(key=lambda x: x["rank"])
    missing = [k for k in ("s1_title", "s2_title", "s3_title") if not content.get(k)]
    if missing:
        sys.exit(f"Slide_Content is missing: {', '.join(missing)}")
    return content, rows


# ---------------------------------------------------------------- helpers
def rgb(hex_):
    return RGBColor.from_string(hex_)


def style_run(run, size, color=WHITE, bold=False, italic=False, font=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    run.font.name = font or (BOLD_FONT if bold else LIGHT)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if name:
        tb.name = name
    return tf


def add_para(tf, runs, first=False, bullet=False, space_before=0, space_after=0, align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    for text, kw in runs:
        r = p.add_run()
        r.text = text
        style_run(r, **kw)
    if space_before:
        p.space_before = Pt(space_before)
    if space_after:
        p.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if bullet:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.2))))
        pPr.set("indent", str(-Emu(Inches(0.2))))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
    return p


def box(slide, x, y, w, h, fill=CARD, rounded=True, line=LINE, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def cell_border(cell, color=LINE, width=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.makeelement(qn(tag), {"w": str(width)})
        sf = ln.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(sf)
        tcPr.append(ln)


def set_cell(cell, text, fill, color=WHITE, bold=False, size=11, align=PP_ALIGN.LEFT):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    style_run(r, size, color, bold)
    cell_border(cell)


def notes(slide, text):
    lines = [l for l in (text or "").splitlines() if l.strip()]
    slide.notes_slide.notes_text_frame.text = "\n".join(["[Sources]", *lines, "[/Sources]"])


# ---------------------------------------------------------------- template handling
def find_slides(prs):
    cover = content = None
    for s in prs.slides:
        texts = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if "Insert Slide Title" in texts and content is None:
            content = s
        elif cover is None and any(sh.is_placeholder and sh.placeholder_format.type == 1 for sh in s.shapes):
            cover = s
    if content is None:
        sys.exit('Template content slide not found (needs a text box reading "Insert Slide Title").')
    return cover, content


def clone_slide(prs, base):
    new = prs.slides.add_slide(base.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in base.shapes:
        new.shapes._spTree.insert_element_before(copy.deepcopy(shp._element), "p:extLst")
    return new


def set_heading(slide, title, client_line):
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        txt = shp.text_frame.text.strip()
        if txt in ("Insert Slide Title", "Client Name"):
            runs = shp.text_frame.paragraphs[0].runs
            runs[0].text = title if txt == "Insert Slide Title" else client_line
            for r in runs[1:]:
                r.text = ""
            if txt == "Insert Slide Title":
                shp.width = Inches(12.1)
                if len(title) > 70:  # keep long action titles on one line
                    for r in runs:
                        r.font.size = Pt(22 if len(title) <= 80 else 20)
                if len(title) > 85:
                    print(f"Warning: title over 85 characters may wrap: {title}")


def set_cover(cover, title, date):
    for shp in cover.shapes:
        if not (shp.is_placeholder and shp.has_text_frame):
            continue
        value = title if shp.placeholder_format.type == 1 else date
        if value:
            runs = shp.text_frame.paragraphs[0].runs
            if runs:
                runs[0].text = value
                for r in runs[1:]:
                    r.text = ""
            else:
                shp.text_frame.text = value


# ---------------------------------------------------------------- slides
def slide_summary(s, c):
    for i in range(3):
        y = 1.4 + i * 1.72
        box(s, LEFT, y, 3.5, 1.52, name=f"stat-card-{i + 1}")
        tf = textbox(s, LEFT + 0.25, y + 0.18, 3.0, 0.62)
        add_para(tf, [(c.get(f"stat{i + 1}", ""), dict(size=32, color=RULE[i], bold=True))], first=True)
        tf = textbox(s, LEFT + 0.25, y + 0.88, 3.0, 0.55)
        add_para(tf, [(c.get(f"stat{i + 1}_text", ""), dict(size=11.5))], first=True)
    tf = textbox(s, 4.45, 1.35, RIGHT - 4.45, 5.0, name="summary-body")
    first = True
    for label, keys in (("Context", ["s1_context"]), ("Key findings", ["s1_finding1", "s1_finding2", "s1_finding3"]),
                        ("Top priorities (score out of 5)", ["s1_priorities"]), ("Implication", ["s1_implication"])):
        add_para(tf, [(label, dict(size=13, color=YELLOW, bold=True))], first=first, space_before=0 if first else 7)
        first = False
        for k in keys:
            if c.get(k):
                add_para(tf, [(c[k], dict(size=12.5))], bullet=True, space_after=3)
    tf = textbox(s, LEFT, 6.58, WIDTH, 0.25)
    add_para(tf, [(c.get("s1_footnote", ""), dict(size=9, color=GREY, italic=True))], first=True)


def slide_priorities(s, c, rows):
    top = [r for r in rows if r["tier"] in ("High", "Medium")][:12]
    low = [r for r in rows if r not in top]
    widths = [0.4, 2.62, 1.45, 1.75, 0.62, 0.8, 3.82]
    hdr = ["#", "Sub-theme", "NBSAP $ | no.", "Client exposure", "Score", "Tier", "Key client evidence"]
    row_h = 0.39 if len(top) <= 9 else 0.33
    gf = s.shapes.add_table(len(top) + 1, 7, Inches(LEFT), Inches(1.3), Inches(WIDTH), Inches(0.36 + row_h * len(top)))
    gf.name = "priority-table"
    tbl = gf.table
    tbl.first_row = tbl.horz_banding = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(0.36)
    for i, h in enumerate(hdr):
        set_cell(tbl.cell(0, i), h, CARD, YELLOW, True, align=PP_ALIGN.CENTER if i in (0, 4, 5) else PP_ALIGN.LEFT)
    for n, r in enumerate(top, 1):
        tbl.rows[n].height = Inches(row_h)
        f = ROW1 if n % 2 else ROW2
        vals = [str(r["rank"]), r["name"], r["nbsap"], r["exposure"], f"{r['score']:.1f}", r["tier"], r["evidence"]]
        for i, v in enumerate(vals):
            if i == 5:
                set_cell(tbl.cell(n, i), v, YELLOW if v == "High" else MEDIUM, DARK if v == "High" else WHITE, True, align=PP_ALIGN.CENTER)
            else:
                set_cell(tbl.cell(n, i), v, f, YELLOW if i == 4 else WHITE, i in (1, 4), align=PP_ALIGN.CENTER if i in (0, 4) else PP_ALIGN.LEFT)
    y = 1.3 + 0.36 + row_h * len(top) + 0.3
    low_txt = "  |  ".join(f"{r['name']} {r['score']:.1f}" for r in low) or "None"
    ph = 0.72 if len(low_txt) + len(c.get("s2_low_note", "")) < 230 else 0.9
    box(s, LEFT, y, WIDTH, ph, name="low-panel")
    tf = textbox(s, LEFT + 0.2, y + 0.05, WIDTH - 0.4, ph - 0.1, MSO_ANCHOR.MIDDLE)
    add_para(tf, [("Low (below 3.0): ", dict(size=11, color=YELLOW, bold=True)), (low_txt + ". ", dict(size=11)),
                  (c.get("s2_low_note", ""), dict(size=11, color=GREY, italic=True))], first=True)
    tf = textbox(s, LEFT, 6.45, WIDTH, 0.42)
    add_para(tf, [(c.get("s2_footnote", ""), dict(size=9, color=GREY, italic=True))], first=True)


def slide_assist(s, c):
    cw, gap, y = 2.69, 0.233, 1.35
    for i in range(4):
        x = LEFT + i * (cw + gap)
        k = f"s3_card{i + 1}_"
        box(s, x, y, cw, 3.82, name=f"card-{i + 1}")
        box(s, x + 0.2, y + 0.2, 0.45, 0.45, fill=RULE[i], rounded=False, line=None)
        tf = textbox(s, x + 0.2, y + 0.2, 0.45, 0.45, MSO_ANCHOR.MIDDLE)
        add_para(tf, [(str(i + 1), dict(size=16, color=DARK, bold=True))], first=True, align=PP_ALIGN.CENTER)
        tf = textbox(s, x + 0.8, y + 0.17, cw - 0.95, 0.7)
        add_para(tf, [(c.get(k + "title", ""), dict(size=13.5, bold=True))], first=True)
        add_para(tf, [("Score " + c.get(k + "score", ""), dict(size=10.5, color=RULE[i]))])
        tf = textbox(s, x + 0.18, y + 0.98, cw - 0.36, 2.78)
        first = True
        for label, key in (("What the public evidence shows", "evidence"), ("How EY could support", "support"), ("Indicative output", "output")):
            add_para(tf, [(label, dict(size=10.5, color=YELLOW, bold=True))], first=first)
            add_para(tf, [(c.get(k + key, ""), dict(size=11.5))], space_after=7)
            first = False
    tf = textbox(s, LEFT, 5.3, WIDTH, 0.3)
    add_para(tf, [("Across all four: ", dict(size=12, color=YELLOW, bold=True)), (c.get("s3_crosscutting", ""), dict(size=12))], first=True)
    box(s, LEFT, 5.72, WIDTH, 0.72, name="next-steps-panel")
    tf = textbox(s, LEFT + 0.22, 5.76, WIDTH - 0.44, 0.64, MSO_ANCHOR.MIDDLE)
    steps = "   ".join(f"{n}. {c[k]}" for n, k in enumerate(("s3_next1", "s3_next2", "s3_next3"), 1) if c.get(k))
    add_para(tf, [("Proposed next steps   ", dict(size=12, color=YELLOW, bold=True)), (steps, dict(size=12))], first=True)
    tf = textbox(s, LEFT, 6.5, WIDTH, 0.25)
    add_para(tf, [(c.get("s3_footnote", ""), dict(size=9, color=GREY, italic=True))], first=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", required=True)
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    c, rows = read_workbook(a.workbook)
    prs = Presentation(a.template)
    cover, base = find_slides(prs)
    s2, s3 = clone_slide(prs, base), clone_slide(prs, base)
    client = c.get("client_name", "")
    if cover is not None:
        set_cover(cover, c.get("cover_title", ""), c.get("cover_date", ""))
    for slide, key in ((base, "s1"), (s2, "s2"), (s3, "s3")):
        set_heading(slide, c[f"{key}_title"], f"{client}  |  {c.get(key + '_section', '')}")
        notes(slide, c.get(f"{key}_notes", ""))
    slide_summary(base, c)
    slide_priorities(s2, c, rows)
    slide_assist(s3, c)
    prs.save(a.out)
    print(f"Saved {a.out}")


if __name__ == "__main__":
    main()
